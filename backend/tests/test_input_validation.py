"""입력 검증·인젝션 방어 회귀 테스트.

SQL 인젝션은 ORM 바인딩으로 이미 막혀 있어 "페이로드가 데이터로만 취급된다"를 확인하고,
PDF 마크업 주입·엑셀 수식 주입·LIKE 와일드카드·범위 초과 입력(500 오류)처럼 실제로 뚫려 있던
부분은 고친 동작을 고정한다.
"""

import io
import zipfile

import openpyxl
import pytest

from app.limits import (
    BCRYPT_MAX_BYTES,
    EMPLOYEE_NAME_MAX_LENGTH,
    MAX_CHILDREN_8_TO_20,
    MAX_DEPENDENTS,
    MAX_PAGE,
    MAX_PAY_AMOUNT,
    MAX_RECORD_ID,
    SEARCH_TEXT_MAX_LENGTH,
)
from tests.conftest import TEST_USERNAME


def _calculate(client, **overrides):
    payload = {"gross_pay": 3_000_000, "num_dependents": 1, **overrides}
    return client.post("/calculate", json=payload)


def _register(client, **overrides):
    payload = {
        "username": "validuser",
        "password": "abcd1234",
        "security_question": "질문",
        "security_answer": "답변",
        **overrides,
    }
    return client.post("/auth/register", json=payload)


# ---------- 급여 입력 범위 ----------


@pytest.mark.parametrize(
    "overrides",
    [
        {"gross_pay": -1},
        {"gross_pay": MAX_PAY_AMOUNT + 1},
        {"gross_pay": 10**30},
        {"bonus_pay": -1},
        {"bonus_pay": MAX_PAY_AMOUNT + 1},
        {"num_dependents": 0},
        {"num_dependents": MAX_DEPENDENTS + 1},
        {"num_children_8_to_20": -1},
        {"num_children_8_to_20": MAX_CHILDREN_8_TO_20 + 1},
        {"employee_name": "가" * (EMPLOYEE_NAME_MAX_LENGTH + 1)},
        {"employee_name": "홍길동\n김철수"},
        {"employee_name": "홍길동\x00"},
    ],
)
def test_calculate_rejects_out_of_range_input(client, overrides):
    assert _calculate(client, **overrides).status_code == 422


def test_calculate_accepts_values_exactly_at_the_limits(client):
    response = _calculate(
        client,
        gross_pay=MAX_PAY_AMOUNT,
        bonus_pay=MAX_PAY_AMOUNT,
        num_dependents=MAX_DEPENDENTS,
        num_children_8_to_20=MAX_CHILDREN_8_TO_20,
        employee_name="가" * EMPLOYEE_NAME_MAX_LENGTH,
    )

    assert response.status_code == 200
    # 세전 급여 + 상여금이 PostgreSQL Integer(int4) 범위 안에 들어와야 운영 DB에서도 저장된다.
    assert response.json()["gross_pay"] + response.json()["bonus_pay"] < 2**31


def test_update_record_rejects_out_of_range_input(client):
    created = _calculate(client).json()

    response = client.put(f"/records/{created['id']}", json={"gross_pay": -5})

    assert response.status_code == 422


def test_employee_name_is_trimmed_and_html_is_stored_as_plain_text(client):
    response = _calculate(client, employee_name="  <script>alert(1)</script>  ")

    # 화면은 React가 텍스트로 렌더링하므로 저장 시에는 변형하지 않고 공백만 정리한다.
    assert response.status_code == 200
    assert response.json()["employee_name"] == "<script>alert(1)</script>"


# ---------- PDF 마크업 주입 ----------


@pytest.mark.parametrize("name", ["<b>굵게", "<img src='x'/>", "a & b < c", "</para>", "<font size=900>"])
def test_payslip_pdf_treats_employee_name_as_plain_text(client, name):
    created = _calculate(client, employee_name=name).json()

    response = client.get(f"/records/{created['id']}/payslip")

    assert response.status_code == 200
    assert response.content.startswith(b"%PDF")


def test_one_markup_name_does_not_break_the_whole_payslip_zip(client):
    _calculate(client, employee_name="<b>굵게")
    _calculate(client, employee_name="홍길동")

    response = client.get("/records/payslips")

    assert response.status_code == 200
    with zipfile.ZipFile(io.BytesIO(response.content)) as zf:
        assert len(zf.namelist()) == 2


# ---------- 엑셀 수식 주입 ----------


@pytest.mark.parametrize("name", ["=1+1", "=HYPERLINK(\"http://evil.example\",\"click\")", "+cmd", "@SUM(A1)"])
def test_excel_export_never_writes_employee_name_as_a_formula(client, name):
    _calculate(client, employee_name=name)

    response = client.get("/records/export")

    sheet = openpyxl.load_workbook(io.BytesIO(response.content)).active
    name_cell = sheet.cell(row=2, column=1)
    assert name_cell.value == name
    assert name_cell.data_type != "f"


def test_excel_export_keeps_numbers_as_numbers(client):
    _calculate(client, employee_name="=1+1", gross_pay=3_000_000)

    sheet = openpyxl.load_workbook(io.BytesIO(client.get("/records/export").content)).active

    assert sheet.cell(row=2, column=2).value == 3_000_000
    assert sheet.cell(row=2, column=2).data_type == "n"


# ---------- 검색/필터: SQL 인젝션·LIKE 와일드카드 ----------


@pytest.mark.parametrize(
    "payload",
    [
        "' OR '1'='1",
        "'; DROP TABLE salary_records; --",
        "\" OR 1=1 --",
        "%' UNION SELECT hashed_password FROM users --",
    ],
)
def test_sql_injection_payload_in_filters_is_treated_as_data(client, payload):
    _calculate(client, employee_name="홍길동")

    for params in ({"search": payload}, {"employee_name": payload}):
        response = client.get("/records", params=params)
        assert response.status_code == 200
        assert response.json()["total"] == 0

    # 인젝션이 실제로 실행됐다면 테이블이 사라지거나 행이 노출됐을 것이다.
    assert client.get("/records").json()["total"] == 1


def test_search_treats_like_wildcards_as_literal_characters(client):
    _calculate(client, employee_name="ab_c")
    _calculate(client, employee_name="abXc")
    _calculate(client, employee_name="50%할인")
    _calculate(client, employee_name="홍길동")

    def names(search):
        items = client.get("/records", params={"search": search}).json()["items"]
        return sorted(item["employee_name"] for item in items)

    assert names("b_c") == ["ab_c"]
    assert names("_") == ["ab_c"]
    assert names("%") == ["50%할인"]
    assert names("\\") == []


def test_search_is_still_a_case_insensitive_substring_match(client):
    _calculate(client, employee_name="Hong Gildong")

    items = client.get("/records", params={"search": "gil"}).json()["items"]

    assert [item["employee_name"] for item in items] == ["Hong Gildong"]


@pytest.mark.parametrize(
    "params",
    [
        {"start_date": "not-a-date"},
        {"end_date": "2026-13-45"},
        {"start_date": "'; DROP TABLE salary_records; --"},
        {"min_gross_pay": MAX_PAY_AMOUNT + 1},
        {"max_gross_pay": -1},
        {"min_gross_pay": 10**20},
        {"search": "가" * (SEARCH_TEXT_MAX_LENGTH + 1)},
        {"employee_name": "가" * (SEARCH_TEXT_MAX_LENGTH + 1)},
    ],
)
@pytest.mark.parametrize("path", ["/records", "/records/export", "/records/payslips"])
def test_malformed_filter_params_are_rejected_with_422_not_500(client, path, params):
    assert client.get(path, params=params).status_code == 422


@pytest.mark.parametrize(
    "params",
    [
        {"page": MAX_PAGE + 1},
        {"page": 10**20},
        {"page": 0},
        {"page_size": 101},
        {"sort_order": "<script>"},
    ],
)
def test_malformed_list_params_are_rejected_with_422_not_500(client, params):
    assert client.get("/records", params=params).status_code == 422


def test_valid_date_filters_still_work(client):
    _calculate(client)

    response = client.get(
        "/records", params={"start_date": "2000-01-01", "end_date": "2100-01-01T23:59:59"}
    )

    assert response.status_code == 200
    assert response.json()["total"] == 1


@pytest.mark.parametrize("method", ["get", "put", "delete"])
def test_record_id_beyond_integer_range_is_rejected_not_a_server_error(client, method):
    path = f"/records/{MAX_RECORD_ID + 1}"
    if method == "get":
        path += "/payslip"

    kwargs = {"json": {"gross_pay": 1}} if method == "put" else {}
    assert getattr(client, method)(path, **kwargs).status_code == 422


# ---------- 회원가입/로그인 ----------


@pytest.mark.parametrize(
    "username",
    ["ab", "a" * 31, "<script>alert(1)</script>", "user name", "user@mail", "a'; DROP TABLE users;--", "한글\n줄바꿈"],
)
def test_register_rejects_invalid_usernames(client, username):
    response = _register(client, username=username)

    assert response.status_code == 422
    assert "아이디" in response.json()["detail"][0]["msg"]


@pytest.mark.parametrize("username", ["abc", "a" * 30, "minji_123", "kim.minji-1", "홍길동01"])
def test_register_accepts_valid_usernames(client, username):
    assert _register(client, username=username).status_code == 200


def test_register_rejects_password_longer_than_bcrypt_limit(client):
    # 영문 73자, 그리고 글자 수는 적지만 UTF-8로 73바이트가 넘는 한글 25자(75바이트)
    for password in ("a1" + "b" * (BCRYPT_MAX_BYTES - 1), "a1" + "가" * 25):
        response = _register(client, password=password)

        assert response.status_code == 422
        assert "너무 깁니다" in response.json()["detail"][0]["msg"]


def test_register_accepts_password_exactly_at_bcrypt_limit(client):
    password = "a1" + "b" * (BCRYPT_MAX_BYTES - 2)
    assert len(password.encode()) == BCRYPT_MAX_BYTES

    assert _register(client, password=password).status_code == 200
    login = client.post("/auth/login", json={"username": "validuser", "password": password})
    assert login.status_code == 200


def test_register_rejects_overlong_security_fields(client):
    assert _register(client, security_question="가" * 101).status_code == 422
    assert _register(client, security_answer="가" * 25).status_code == 422


def test_login_with_overlong_password_is_unauthorized_not_a_server_error(client):
    response = client.post("/auth/login", json={"username": TEST_USERNAME, "password": "a" * 100})

    assert response.status_code == 401


def test_login_rejects_absurdly_long_credentials_before_hashing(client):
    response = client.post("/auth/login", json={"username": "u" * 5000, "password": "p" * 5000})

    assert response.status_code == 422


def test_login_sql_injection_payload_is_just_a_wrong_username(client):
    response = client.post(
        "/auth/login", json={"username": "' OR '1'='1' --", "password": "whatever1"}
    )

    assert response.status_code == 401


def test_password_reset_with_overlong_security_answer_fails_cleanly(client):
    response = client.post(
        "/auth/password-reset",
        json={"username": TEST_USERNAME, "security_answer": "가" * 50, "new_password": "newpass123"},
    )

    assert response.status_code == 400


def test_change_password_rejects_overlong_new_password(client):
    response = client.put(
        "/auth/password",
        json={"current_password": "testpass123", "new_password": "a1" + "b" * 80},
    )

    assert response.status_code == 422


# ---------- CSV 일괄 업로드 ----------


def _upload(client, csv_text: str):
    return client.post(
        "/calculate/bulk", files={"file": ("a.csv", csv_text.encode("utf-8"), "text/csv")}
    )


def test_csv_rows_with_out_of_range_values_are_reported_and_valid_rows_still_saved(client):
    csv_text = (
        "employee_name,gross_pay,bonus_pay,num_dependents,num_children_8_to_20\n"
        "정상,3000000,0,1,0\n"
        f"급여초과,{MAX_PAY_AMOUNT + 1},0,1,0\n"
        "무한대,inf,0,1,0\n"
        "지수,1e30,0,1,0\n"
        f"상여초과,3000000,{MAX_PAY_AMOUNT + 1},1,0\n"
        f"부양초과,3000000,0,{MAX_DEPENDENTS + 1},0\n"
        f"자녀초과,3000000,0,1,{MAX_CHILDREN_8_TO_20 + 1}\n"
    )

    response = _upload(client, csv_text)

    assert response.status_code == 200
    body = response.json()
    assert [r["employee_name"] for r in body["created"]] == ["정상"]
    assert [e["row"] for e in body["errors"]] == [3, 4, 5, 6, 7, 8]
    assert client.get("/records").json()["total"] == 1


def test_csv_rows_with_invalid_employee_names_are_reported(client):
    too_long = "가" * (EMPLOYEE_NAME_MAX_LENGTH + 1)
    csv_text = f'employee_name,gross_pay\n"줄\n바꿈",3000000\n{too_long},3000000\n정상,3000000\n'

    body = _upload(client, csv_text).json()

    assert [r["employee_name"] for r in body["created"]] == ["정상"]
    assert [e["row"] for e in body["errors"]] == [2, 3]
    assert all("직원명" in e["reason"] for e in body["errors"])


def test_csv_numeric_and_empty_employee_names_are_stored_as_text(client):
    body = _upload(client, "employee_name,gross_pay\n1234,3000000\n,3000000\n").json()

    # 숫자처럼 보이는 이름은 문자열로, 빈 칸은 "nan"이 아니라 빈 문자열로 저장된다.
    assert [r["employee_name"] for r in body["created"]] == ["1234", ""]


def test_csv_formula_like_employee_name_is_exported_as_text(client):
    _upload(client, "employee_name,gross_pay\n=1+1,3000000\n")

    sheet = openpyxl.load_workbook(io.BytesIO(client.get("/records/export").content)).active

    assert sheet.cell(row=2, column=1).value == "=1+1"
    assert sheet.cell(row=2, column=1).data_type != "f"
